"""Медийка Ozon: недельная таблица расходов в формате общей таблицы затрат.

usage: python3 -I build_week.py <days.csv> <out.xlsx> [export_1.xlsx export_2.xlsx]
days.csv — дата;каб1;каб2 (расход за день по каждому кабинету, ₽), например 19.10.2026;4600,42;3336,74
Выгрузки «Статистика по кампаниям» за ту же неделю — для сверки итога (необязательно).

Паттерн (см. README.md): дневной расход обоих кабинетов ÷ 5 → одинаковая сумма на каждый из 5 артикулов,
сумма с копейками (итого дня ÷ 5), по блоку строк на артикул.
"""
import csv, datetime, sys, os
sys.path.insert(0, os.environ.get("PYLIB", ""))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ARTICLES = ["SHVPAREKCHB", "SHVPAR7V1CH", "VACCLCOMPACT", "VACCLDRY", "AEROGRIL2CHASHI"]  # порядок как в общей таблице
DAYS, OUT = sys.argv[1:3]
EXPORTS = sys.argv[3:]


def num(s):
    return float(str(s).replace(" ", "").replace(" ", "").replace(",", "."))


days = []
for row in csv.reader(open(DAYS, encoding="utf-8"), delimiter=";"):
    if not row or row[0].startswith("#"):
        continue
    days.append((datetime.datetime.strptime(row[0].strip(), "%d.%m.%Y"), num(row[1]), num(row[2])))

wb = openpyxl.Workbook()
B = Font(bold=True)
HDR = PatternFill("solid", fgColor="F4F2EE")

# ── лист 1: расход по дням
d = wb.active
d.title = "Расход по дням"
d.append(["Дата", "Каб 1, ₽", "Каб 2, ₽", "Итого, ₽", "÷5, ₽"])
for i, (dt, k1, k2) in enumerate(days, 2):
    d.append([dt, k1, k2, f"=B{i}+C{i}", f"=D{i}/5"])
last = len(days) + 1
d.append(["Итого", f"=SUM(B2:B{last})", f"=SUM(C2:C{last})", f"=SUM(D2:D{last})", f"=SUM(E2:E{last})"])
tot = last + 1
if EXPORTS:
    ex = []
    for f in EXPORTS:
        ws = openpyxl.load_workbook(f, data_only=True).worksheets[0]
        ex.append(round(sum(r[12] or 0 for r in ws.iter_rows(min_row=2, values_only=True) if r[0]), 2))
    d.append([])
    d.append(["Сверка с выгрузками (итог за неделю)", *ex, f"=SUM(B{tot + 2}:C{tot + 2})"])
    d.append(["Разница", f"=B{tot}-B{tot + 2}", f"=C{tot}-C{tot + 2}", f"=D{tot}-D{tot + 2}"])
for c in d[1]:
    c.font = B; c.fill = HDR
for c in d[tot]:
    c.font = B
for r in d.iter_rows(min_row=2):
    r[0].number_format = "DD.MM.YYYY"
    for c in r[1:]:
        c.number_format = "# ##0.00"
for col, w in zip("ABCDE", (34, 14, 14, 14, 12)):
    d.column_dimensions[col].width = w

# ── лист 2: строки для общей таблицы (как в примере)
m = wb.create_sheet("Медийка")
n = len(days) * len(ARTICLES)
m["C1"] = f"=SUBTOTAL(9,C3:C{n + 2})"
m.append(["Артикул", "Дата рекламы", "Затраты на рекламу c НДС", "Затраты на рекламу б/НДС", "Площадка маркетплейс",
          "Комментарии \n(где реклама)", "Неделя", "Месяц"])
r = 3
for art in ARTICLES:
    for i, (dt, k1, k2) in enumerate(days):
        m.append([art, dt, round((k1 + k2) / 5, 2), f'=IF(F{r}= "директ", C{r}/1.22, C{r})', "Озон", "Медийная реклама",
                  f"=ISOWEEKNUM(B{r})", f"=MONTH(B{r})"])
        r += 1
for c in m[2]:
    c.font = B; c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical="top")
for row in m.iter_rows(min_row=3):
    row[1].number_format = "DD.MM.YYYY"
    row[2].number_format = row[3].number_format = "# ##0.00"
for col, w in zip("ABCDEFGH", (18, 13, 14, 14, 12, 18, 8, 8)):
    m.column_dimensions[col].width = w
m.freeze_panes = "A3"
wb.active = 1
wb.save(OUT)
tot_days = sum(k1 + k2 for _, k1, k2 in days)
print("ok", OUT, "дней", len(days), "итого", round(tot_days, 2), "÷5", round(tot_days / 5, 2),
      "строк", n, "сумма строк", sum(round((k1 + k2) / 5, 2) for _, k1, k2 in days) * len(ARTICLES))
