"""Сборка Excel + YML для Яндекс Директа.
usage: python3 build.py <products.json> <semantika.csv> <out_dir>
products.json — выжимка карточек luxhomme.store (extract.py), semantika.csv — research/semantika-vertikalnyj-moyushhij-pylesos.csv
"""
import csv, datetime, json, os, re, sys
from xml.sax.saxutils import escape

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

PJ, SEM, OUT = sys.argv[1:4]
D = json.load(open(PJ, encoding="utf-8"))
BY = {p["slug"]: p for p in C.P}
LIM = {"h1": 56, "h2": 30, "txt": 81, "disp": 20, "sl_t": 30, "sl_d": 60, "co": 25, "notes": 50}
errors = []


def money(x):
    return f"{x:,.0f}".replace(",", " ")


def fill(s, p):
    return s.replace("{PROMO}", C.PROMO).replace("{P20}", money(p["p20"])).replace("{PRICE}", money(p["price"]))


def chk(kind, s, where):
    if len(s) > LIM[kind]:
        errors.append(f"{where}: {kind} {len(s)}>{LIM[kind]}: {s}")
    return s


# ── цены, ключи, ссылки
for p in C.P:
    d = D[p["slug"]]
    p["name"], p["sku"], p["price"] = d["name"], d["sku"] or "", int(float(d["price"]))
    p["p20"] = round(p["price"] * (1 - C.DISCOUNT))
    p["url"] = d["url"]
    p["image"] = d["image"]
    p["ld"] = d["ld_description"]
    p["avail"] = "InStock" in (d["availability"] or "")
    p["key"] = p["disp"].lower()


def utm(p, kind, path=None, content_suffix=""):
    camp = {"search": f"lux_search_{p['cat']}", "epk": f"lux_epk_{p['cat']}", "feed": "lux_feed"}[kind]
    base = C.SITE + path if path else p["url"]
    return (f"{base}?utm_source=yandex&utm_medium=cpc&utm_campaign={camp}_{{campaign_id}}"
            f"&utm_content={p['key']}{content_suffix}_{{ad_id}}_{{source_type}}&utm_term={{keyword}}")


# ── семантика Aqvion из файла клиента
def stem_hit(word, minus):
    return word == minus or (len(minus) >= 5 and word.startswith(minus[:-2]))


def blocked(phrase, minus_list):
    ws = phrase.split()
    for m in minus_list:
        mw = m.split()
        if all(any(stem_hit(w, x) for w in ws) for x in mw):
            return True
    return False


def minus_for(p):
    out = []
    for v in C.NEG_COMMON.values():
        out += v.split(", ")
    out += C.NEG_COMP[p["cat"]].split(", ") + p["neg"].split(", ")
    return [x.strip() for x in out if x.strip()]


for p in C.P:
    p["freq"] = {}
    if p.get("kw_csv"):
        rows = list(csv.reader(open(SEM, encoding="utf-8")))[1:]
        mins = minus_for(p) + ["2026", "выбирать", "как", "самый", "хороший", "лучше",
                              # нет в карточке Aqvion Pro или мусорные хвосты
                              "паровые", "паром", "фирмы", "red", "history", "h2", "стоит", "pro", "про", "белый",
                              "ручка", "180", "5", "насадками", "легкий"]
        exact = {"пылесосит"}
        groups = {"Тип + покупка": [], "Самоочистка и база": [], "Беспроводной": [], "Назначение": [], "Тип": []}
        n = 0
        for q, f in rows:
            q = q.strip().lower()
            if int(f) < 500:  # хвост ниже 500 — в основном чужие модели
                break
            if not re.search(r"моющ|моет", q) or blocked(q, mins) or exact & set(q.split()):
                continue
            g = ("Тип + покупка" if re.search(r"купить|цена|цены|недорог|бюджет", q) else
                 "Самоочистка и база" if re.search(r"самоочист|сушк|станци|базой", q) else
                 "Беспроводной" if re.search(r"беспровод|аккумулятор", q) else
                 "Назначение" if re.search(r"для дома|квартир|полов|ковр|сухая|сухой|влажн", q) else "Тип")
            groups[g].append(q)
            p["freq"][q] = int(f)
            n += 1
            if n >= 45:
                break
        groups.update(p["kw"])
        p["kw"] = {k: v for k, v in groups.items() if v}
    for g, lst in p["kw"].items():
        for q in lst:
            if blocked(q, minus_for(p)):
                errors.append(f"{p['short']}: фраза «{q}» режется минус-словами")

# ── стили
B = Font(bold=True)
H1F = Font(bold=True, size=14)
SEC = PatternFill("solid", fgColor="EDE6DD")
HDR = PatternFill("solid", fgColor="F6F2EC")
thin = Side(style="thin", color="D9D2C7")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def put(ws, r, vals, bold=False, fillc=None, box=True):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.alignment = WRAP
        if bold:
            c.font = B
        if fillc:
            c.fill = fillc
        if box and v is not None:
            c.border = BOX
    return r + 1


def section(ws, r, title, ncols=8):
    for i in range(1, ncols + 1):
        ws.cell(row=r, column=i).fill = SEC
    ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12)
    return r + 1


def widths(ws, ws_list):
    for i, w in enumerate(ws_list, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def L(col, row):
    return f"=LEN({col}{row})"


# ── 1. Обзор
ws = wb.active
ws.title = "Обзор"
widths(ws, [30, 30, 22, 34, 40, 46])
r = 1
ws.cell(row=r, column=1, value="Luxhommè — Яндекс Директ на сайт: семантика, объявления, UTM (15 товаров)").font = H1F
r += 2
r = section(ws, r, "Параметры", 6)
for a, b in [("Сайт", C.SITE), ("Промокод", C.PROMO), ("Скидка", "20% на все товары при заказе на сайте"),
             ("География", "Вся Россия"), ("Кампании", "Поиск по товарам · Сети (РСЯ, ЕПК) · Товарная по фиду"),
             ("Факты о товарах", "Карточки luxhomme.store, снято 09.10.2026"),
             ("Частотности", "Только Aqvion Pro — из файла семантики клиента (research/semantika-vertikalnyj-moyushhij-pylesos.csv). По остальным не снимались"),
             ("Конкуренты", "Запросы с брендами конкурентов — в минус-словах"),
             ("Бренд Luxhommè", "В минус-словах: брендовые запросы идут в действующую брендовую кампанию")]:
    r = put(ws, r, [a, b], bold=False)
    ws.cell(row=r - 1, column=1).font = B
r += 1
r = section(ws, r, "Структура кампаний", 6)
r = put(ws, r, ["Кампания", "Тип / площадки", "Группы", "utm_campaign", "Стратегия на старте", "Где в файле"], bold=True, fillc=HDR)
for cat, nm in C.CATS.items():
    items = [BY[s]["short"] for s in C.ORDER if BY[s]["cat"] == cat]
    r = put(ws, r, [f"Lux | Поиск | {nm}", "ЕПК, только поиск", ", ".join(items), f"lux_search_{cat}_{{campaign_id}}",
                    "Максимум кликов с недельным бюджетом → Максимум конверсий по цели «Заказ» после накопления конверсий", "Листы товаров, «Импорт Поиск»"])
for cat, nm in C.CATS.items():
    items = [BY[s]["short"] for s in C.ORDER if BY[s]["cat"] == cat]
    r = put(ws, r, [f"Lux | Сети | {nm}", "ЕПК, только РСЯ: автотаргетинг + фразы, картинки", ", ".join(items), f"lux_epk_{cat}_{{campaign_id}}",
                    "Как в поиске", "Листы товаров (блок РСЯ), «Импорт Сети»"])
r = put(ws, r, ["Lux | Товарная | Фид", "Товарная кампания: поиск + РСЯ", "15 офферов из фида", "lux_feed_{campaign_id}",
                "Как в поиске", "«Фид», файл luxhomme-direct-feed.yml"])
r += 1
r = section(ws, r, "UTM-разметка", 6)
r = put(ws, r, ["Параметр", "Значение", "Что даёт"], bold=True, fillc=HDR)
for a, b, c in [("utm_source", "yandex", "Источник"), ("utm_medium", "cpc", "Платный трафик"),
                ("utm_campaign", "lux_<search|epk|feed>_<категория>_{campaign_id}", "Тип кампании, категория, ID кампании"),
                ("utm_content", "<модель>_{ad_id}_{source_type}", "Товар, ID объявления, поиск/сети"),
                ("utm_term", "{keyword}", "Ключевая фраза"),
                ("Быстрые ссылки", "в utm_content добавлено _sitelink", "Отделить переходы по быстрым ссылкам")]:
    r = put(ws, r, [a, b, c])
r += 1
r = section(ws, r, "Перед запуском", 6)
for t in [
    f"Завести на сайте промокод {C.PROMO}: −20% на все товары; указать условие на сайте (баннер или строка в карточке) — модерация сверяет скидку с посадочной.",
    "Гостевое оформление заказа и срок/стоимость доставки до входа в аккаунт (UI/UX-аудит, разделы 1–2).",
    "Гарантия: в карточках 3 года (EcoSteam Pro — 12 месяцев, uFit 2D Standard — 2 года), на /service — 12 месяцев. Привести к одному сроку; в объявлениях срок взят из карточки.",
    "Бесплатную доставку не обещать: на /delivery стоимость не рассчитывается автоматически.",
    "CremaX нет в sitemap — добавить.",
    "На /returns продавец — ИП, в подвале сайта — ООО «САУДАГАР Е-КОМ»: указать одного продавца.",
    "Виброплатформы с EMS: тексты без медицинских обещаний; модерация может запросить документы — подготовить сертификат/декларацию.",
    "Nespresso, Dolce Gusto в текстах — только как совместимые форматы капсул.",
    f"Если модерация отклонит слово заглавными ({C.PROMO}) — заменить на вариант с одной заглавной, код на сайте сделать нечувствительным к регистру.",
    "Метрика: цели «Заказ» и «Применение промокода», счётчик подключён к кампаниям, разметка ссылок для Метрики (yclid) включена.",
    "Колонки листов «Импорт» сверить с актуальным шаблоном Директ Коммандера перед загрузкой.",
]:
    r = put(ws, r, ["☐", t])
    ws.merge_cells(start_row=r - 1, start_column=2, end_row=r - 1, end_column=6)
r += 1
r = section(ws, r, "Листы", 6)
for a, b in [("Товары", "Сводка: цены, цена с промокодом, ссылки с UTM"),
             ("Листы товаров", "УТП, семантика, минус-слова товара, объявления поиска, быстрые ссылки, уточнения, РСЯ, ссылки"),
             ("Минус-слова", "Общие, конкуренты по категориям, по товарам"),
             ("Импорт Поиск / Импорт Сети", "Строки для Директ Коммандера (колонка «Кампания» — фильтр перед вставкой)"),
             ("Фид", "Офферы для товарной кампании; YML — отдельным файлом")]:
    r = put(ws, r, [a, b])
ws.freeze_panes = "A2"

# ── 2. Товары
ws = wb.create_sheet("Товары")
widths(ws, [4, 34, 18, 16, 10, 12, 18, 44, 70, 70, 70])
ws.cell(row=1, column=1, value="Скидка").font = B
ws.cell(row=1, column=2, value=C.DISCOUNT).number_format = "0%"
put(ws, 3, ["№", "Товар", "Категория", "SKU", "Цена, ₽", "С промокодом, ₽", "Гарантия", "Страница", "Ссылка Поиск (UTM)", "Ссылка Сети (UTM)", "Ссылка Фид (UTM)"],
    bold=True, fillc=HDR)
for i, s in enumerate(C.ORDER, 1):
    p = BY[s]
    rr = 3 + i
    put(ws, rr, [i, p["name"], C.CATS[p["cat"]], p["sku"] or "нет в разметке", p["price"], f"=ROUND(E{rr}*(1-$B$1),0)", p["warranty"].replace("Гарантия ", ""),
                 p["url"], utm(p, "search"), utm(p, "epk"), utm(p, "feed")])
    ws.cell(row=rr, column=5).number_format = "# ##0"
    ws.cell(row=rr, column=6).number_format = "# ##0"
ws.freeze_panes = "C4"

# ── 3. Листы товаров
imp_search, imp_net = [], []
for gi, s in enumerate(C.ORDER, 1):
    p = BY[s]
    ws = wb.create_sheet(p["short"][:31])
    widths(ws, [24, 52, 9, 34, 9, 70, 9, 44])
    r = 1
    ws.cell(row=r, column=1, value=p["name"]).font = H1F
    r += 1
    r = put(ws, r, ["Страница", p["url"]])
    r = put(ws, r, ["SKU", p["sku"] or "нет в разметке сайта"])
    r = put(ws, r, ["Цена, ₽", p["price"]])
    r = put(ws, r, ["С промокодом, ₽", f"=ROUND(B{r - 1}*(1-Товары!$B$1),0)"])
    r = put(ws, r, ["Промокод", C.PROMO])
    r = put(ws, r, ["Категория", C.CATS[p["cat"]]])
    r += 1
    r = section(ws, r, "УТП (факты из карточки)")
    for u in p["utp"]:
        r = put(ws, r, ["•", u])
    r += 1
    r = section(ws, r, "Семантика")
    has_f = bool(p["freq"])
    r = put(ws, r, ["Группа", "Фраза"] + (["Частотность"] if has_f else []), bold=True, fillc=HDR)
    for g, lst in p["kw"].items():
        for q in lst:
            r = put(ws, r, [g, q] + ([p["freq"].get(q)] if has_f else []))
    r += 1
    r = section(ws, r, "Минус-слова товара (+ общие и конкуренты — лист «Минус-слова»)")
    r = put(ws, r, ["Товар", p["neg"]])
    ws.merge_cells(start_row=r - 1, start_column=2, end_row=r - 1, end_column=8)
    r = put(ws, r, ["Конкуренты", C.NEG_COMP[p["cat"]]])
    ws.merge_cells(start_row=r - 1, start_column=2, end_row=r - 1, end_column=8)
    r += 1
    r = section(ws, r, "Объявления — Поиск")
    r = put(ws, r, ["№", "Заголовок 1", "≤56", "Заголовок 2", "≤30", "Текст", "≤81", "Отображаемая ссылка (≤20)"], bold=True, fillc=HDR)
    chk("disp", p["disp"], p["short"])
    ads = []
    for i, (h1, h2, t) in enumerate(p["ads"], 1):
        h1, h2, t = fill(h1, p), fill(h2, p), fill(t, p)
        chk("h1", h1, p["short"]); chk("h2", h2, p["short"]); chk("txt", t, p["short"])
        ads.append((h1, h2, t))
        r = put(ws, r, [i, h1, L("B", r), h2, L("D", r), t, L("F", r), p["disp"]])
    r = put(ws, r, ["Ссылка (UTM)", utm(p, "search")])
    ws.merge_cells(start_row=r - 1, start_column=2, end_row=r - 1, end_column=8)
    r += 1
    r = section(ws, r, "Быстрые ссылки")
    r = put(ws, r, ["Название", "Описание", "≤60", "Ссылка (UTM)", "", "", "", "Название ≤30"], bold=True, fillc=HDR)
    sls = list(C.COMMON_SITELINKS)
    if p["sibling"]:
        sib = BY[p["sibling"]]
        sls.append((sib["sl"][0], "/products/" + sib["slug"], sib["sl"][1]))
    else:
        sls.append(("Правила возврата", "/returns", "Возврат товара надлежащего качества в течение 7 дней"))
    sl_rows = []
    for t, path, dsc in sls:
        chk("sl_t", t, p["short"]); chk("sl_d", dsc, p["short"])
        u = utm(p, "search", path, "_sitelink")
        sl_rows.append((t, dsc, u))
        r = put(ws, r, [t, dsc, L("B", r), u, None, None, None, L("A", r)])
        ws.merge_cells(start_row=r - 1, start_column=4, end_row=r - 1, end_column=7)
    r += 1
    r = section(ws, r, "Уточнения")
    callouts = [p["warranty"]] + C.COMMON_CALLOUTS
    r = put(ws, r, ["Уточнение", "≤25"], bold=True, fillc=HDR)
    for c in callouts:
        chk("co", c, p["short"])
        r = put(ws, r, [c, L("A", r)])
    r += 1
    r = section(ws, r, "Объявления — Сети (РСЯ, ЕПК)")
    r = put(ws, r, ["№", "Заголовок", "≤56", "Текст", "≤81", "Бриф на картинку", "", "Фото с сайта"], bold=True, fillc=HDR)
    rs = []
    for i, (h, t) in enumerate(p["rsya"], 1):
        h, t = fill(h, p), fill(t, p)
        chk("h1", h, p["short"] + " РСЯ"); chk("txt", t, p["short"] + " РСЯ")
        rs.append((h, t))
        r = put(ws, r, [i, h, L("B", r), t, L("D", r), f"{p['img']}. Форматы: {C.IMG_FORMATS}", None, p["image"]])
        ws.merge_cells(start_row=r - 1, start_column=6, end_row=r - 1, end_column=7)
    r = put(ws, r, ["Ссылка (UTM)", utm(p, "epk")])
    ws.merge_cells(start_row=r - 1, start_column=2, end_row=r - 1, end_column=8)
    r += 1
    r = section(ws, r, "Ссылки с UTM")
    for a, b in [("Поиск", utm(p, "search")), ("Сети", utm(p, "epk")), ("Фид", utm(p, "feed"))]:
        r = put(ws, r, [a, b])
        ws.merge_cells(start_row=r - 1, start_column=2, end_row=r - 1, end_column=8)

    # строки импорта
    camp_s = f"Lux | Поиск | {C.CATS[p['cat']]}"
    camp_n = f"Lux | Сети | {C.CATS[p['cat']]}"
    slt = "||".join(x[0] for x in sl_rows)
    sld = "||".join(x[1] for x in sl_rows)
    slu = "||".join(x[2] for x in sl_rows)
    co = "||".join(callouts)
    neg_g = p["neg"]
    phrases = [q for lst in p["kw"].values() for q in lst]
    for j, q in enumerate(phrases):
        h1, h2, t = ads[0]
        imp_search.append([camp_s, "-", "Текстово-графическое", "-", p["short"], gi, q, h1, h2, t, utm(p, "search"), p["disp"], "Россия",
                           slt, sld, slu, co, neg_g if j == 0 else ""])
    for h1, h2, t in ads[1:]:
        imp_search.append([camp_s, "+", "Текстово-графическое", "-", p["short"], gi, "", h1, h2, t, utm(p, "search"), p["disp"], "Россия",
                           slt, sld, slu, co, ""])
    net_phr = [q for g, lst in p["kw"].items() if g != "Модель" for q in lst]
    for j, q in enumerate(net_phr):
        h, t = rs[0]
        imp_net.append([camp_n, "-", "Текстово-графическое", "-", p["short"], gi, q, h, "", t, utm(p, "epk"), p["disp"], "Россия",
                        slt, sld, slu, co, neg_g if j == 0 else "", p["image"]])
    for h, t in rs[1:]:
        imp_net.append([camp_n, "+", "Текстово-графическое", "-", p["short"], gi, "", h, "", t, utm(p, "epk"), p["disp"], "Россия",
                        slt, sld, slu, co, "", p["image"]])

# ── 4. Минус-слова
ws = wb.create_sheet("Минус-слова")
widths(ws, [34, 120])
r = 1
r = section(ws, r, "Общие — на все кампании", 2)
for k, v in C.NEG_COMMON.items():
    r = put(ws, r, [k, v])
r += 1
r = section(ws, r, "Конкуренты — на кампании категории", 2)
for k, v in C.NEG_COMP.items():
    r = put(ws, r, [C.CATS[k], v])
r += 1
r = section(ws, r, "По товарам — на группу", 2)
for s in C.ORDER:
    r = put(ws, r, [BY[s]["short"], BY[s]["neg"]])

# ── 5. Импорт
HEAD = ["Кампания", "Доп. объявление группы", "Тип объявления", "Мобильное объявление", "Название группы", "Номер группы",
        "Фраза (с минус-словами)", "Заголовок 1", "Заголовок 2", "Текст", "Ссылка", "Отображаемая ссылка", "Регион",
        "Заголовки быстрых ссылок", "Описания быстрых ссылок", "Адреса быстрых ссылок", "Уточнения", "Минус-слова на группу"]
for title, rows, extra in [("Импорт Поиск", imp_search, []), ("Импорт Сети", imp_net, ["Изображение"])]:
    ws = wb.create_sheet(title)
    put(ws, 1, HEAD + extra, bold=True, fillc=HDR)
    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            ws.cell(row=i, column=j, value=v)
    widths(ws, [26, 8, 18, 8, 20, 7, 44, 50, 26, 70, 60, 18, 10, 40, 60, 60, 50, 50, 50])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEAD + extra))}{len(rows) + 1}"

# ── 6. Фид
ws = wb.create_sheet("Фид")
widths(ws, [16, 44, 18, 10, 30, 70, 50, 44])
notes = f"Скидка 20% по промокоду {C.PROMO} на сайте"
chk("notes", notes, "feed")
put(ws, 1, ["id", "Название", "Категория", "Цена, ₽", "Модель", "Ссылка (UTM)", "Картинка", "sales_notes"], bold=True, fillc=HDR)
cat_id = {k: i for i, k in enumerate(C.CATS, 1)}
offers = []
for i, s in enumerate(C.ORDER, 2):
    p = BY[s]
    oid = p["sku"] or "UFIT3DMAX"
    offers.append((oid, p))
    put(ws, i, [oid, p["name"], C.CATS[p["cat"]], p["price"], p["short"], utm(p, "feed"), p["image"], notes])

# ── YML
now = datetime.datetime(2026, 10, 9, 12, 0).strftime("%Y-%m-%dT%H:%M+03:00")
y = [f'<?xml version="1.0" encoding="UTF-8"?>', f'<yml_catalog date="{now}">', "<shop>", "<name>Luxhommè</name>",
     "<company>ООО «САУДАГАР Е-КОМ»</company>", f"<url>{C.SITE}</url>", '<currencies><currency id="RUR" rate="1"/></currencies>', "<categories>"]
for k, i in cat_id.items():
    y.append(f'<category id="{i}">{escape(C.CATS[k])}</category>')
y += ["</categories>", "<offers>"]
for oid, p in offers:
    desc = re.sub(r"\s+", " ", p["ld"]).strip()[:600]
    y += [f'<offer id="{escape(oid)}" available="{"true" if p["avail"] else "false"}">',
          f"<url>{escape(utm(p, 'feed'))}</url>", f"<price>{p['price']}</price>", "<currencyId>RUR</currencyId>",
          f"<categoryId>{cat_id[p['cat']]}</categoryId>", f"<picture>{escape(p['image'])}</picture>",
          f"<name>{escape(p['name'])}</name>", "<vendor>Luxhommè</vendor>", f"<model>{escape(p['short'])}</model>",
          f"<description>{escape(desc)}</description>", f"<sales_notes>{escape(notes)}</sales_notes>", "</offer>"]
y += ["</offers>", "</shop>", "</yml_catalog>", ""]

os.makedirs(OUT, exist_ok=True)
xl = os.path.join(OUT, "Luxhomme_Yandex-Direct_2026-10.xlsx")
wb.save(xl)
open(os.path.join(OUT, "luxhomme-direct-feed.yml"), "w", encoding="utf-8").write("\n".join(y))
n_kw = sum(len(v) for p in C.P for v in p["kw"].values())
print("ok", xl, "| фраз:", n_kw, "| строк импорта поиск/сети:", len(imp_search), len(imp_net))
print("\n".join(errors) if errors else "длины и минус-слова: ок")
